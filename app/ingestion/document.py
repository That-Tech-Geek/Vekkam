from __future__ import annotations

from pathlib import Path

from app.core.ids import stable_id
from app.core.models import Concept, LearningDocument, PageIR, SourceRegion


def ingest_document(path: Path) -> LearningDocument:
    """Extract text from a PDF or DOCX while preserving page boundaries where possible."""
    path = path.expanduser().resolve()
    suffix = path.suffix.lower()
    document_id = stable_id("document", str(path))
    title = path.stem

    if suffix == ".pdf":
        return _ingest_pdf(path, document_id, title)
    if suffix == ".docx":
        return _ingest_docx(path, document_id, title)
    raise ValueError("Unsupported document type. Use .pdf or .docx.")


def _ingest_pdf(path: Path, document_id: str, title: str) -> LearningDocument:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError(
            "PDF support is not installed. Install the document extra: "
            "python -m pip install -e '.[document]'"
        ) from exc

    pages: list[PageIR] = []
    with fitz.open(path) as pdf:
        for number, page in enumerate(pdf, start=1):
            text = page.get_text("text").strip()
            page_id = stable_id("page", document_id, str(number))
            pages.append(
                PageIR(
                    page_id=page_id,
                    page_number=number,
                    text=text,
                    blocks=[],
                )
            )

    return _document(document_id, title, pages)


def _ingest_docx(path: Path, document_id: str, title: str) -> LearningDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "DOCX support is not installed. Install the document extra: "
            "python -m pip install -e '.[document]'"
        ) from exc

    document = Document(path)
    text = "\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )
    page_id = stable_id("page", document_id, "1")
    pages = [PageIR(page_id=page_id, page_number=1, text=text)]
    return _document(document_id, title, pages)


def _document(document_id: str, title: str, pages: list[PageIR]) -> LearningDocument:
    text = "\n\n".join(page.text for page in pages if page.text.strip())
    page_id = pages[0].page_id if pages else stable_id("page", document_id, "1")
    concept = Concept(
        concept_id=stable_id("concept", document_id, text[:160]),
        name=title,
        description=text[:1000],
        evidence=[
            SourceRegion(page_id=page_id, extracted_text=text[:4000])
        ],
    )
    return LearningDocument(
        document_id=document_id,
        title=title,
        pages=pages,
        concepts=[concept] if text else [],
    )


class DocumentIngestor:
    """Compatibility boundary for document adapters."""

    def ingest(self, path: Path) -> LearningDocument:
        return ingest_document(path)
